import io
import re
from pdfplumber import open as open_pdf
from docx import Document
from utils.normalization import normalize_date, normalize_skill

# Regular expressions for contact extraction
EMAIL_RE = re.compile(r"[a-zA-Z0-9+._%-]+@[a-zA-Z0-9._%-]+\.[a-zA-Z]{2,}")
PHONE_RE = re.compile(r"\+?\d[\d \-]{7,}\d")

def parse_job_description(file_obj) -> str:
    """
    Extract full text from a PDF or DOCX job description file.
    """
    if hasattr(file_obj, "read") and file_obj.name.lower().endswith('.pdf'):
        text = ''
        with open_pdf(file_obj) as pdf:
            for page in pdf.pages:
                text += page.extract_text() or ''
        return text
    elif hasattr(file_obj, "read") and file_obj.name.lower().endswith('.docx'):
        doc = Document(io.BytesIO(file_obj.read()))
        return '\n'.join(p.text for p in doc.paragraphs)
    return ''

def parse_criteria(file_obj) -> str:
    """
    Parse a criteria file (reuse job description parser logic).
    """
    return parse_job_description(file_obj)

def parse_resume(file_obj) -> dict:
    """
    Parse a resume (PDF or DOCX) and extract structured fields.
    Returns a dict with name, email, phone, skills, dates, raw_text.
    """
    text = ''
    # Extract raw text
    if file_obj.name.lower().endswith('.pdf'):
        with open_pdf(file_obj) as pdf:
            for page in pdf.pages:
                text += page.extract_text() or ''
    elif file_obj.name.lower().endswith('.docx'):
        doc = Document(io.BytesIO(file_obj.read()))
        text = '\n'.join(p.text for p in doc.paragraphs)

    # Simple heuristics for field extraction
    name = text.split('\n')[0].strip() if text else None
    email_match = EMAIL_RE.search(text)
    phone_match = PHONE_RE.search(text)

    # Skills: detect a 'Skills' section
    skills = []
    for line in text.splitlines():
        if 'skills' in line.lower():
            parts = re.split(r'[:,]', line, maxsplit=1)
            if len(parts) > 1:
                skills = [normalize_skill(s) for s in parts[1].split(',')]
                break

    # Dates normalization from any 'Month YYYY' patterns
    dates = re.findall(r'\b\w+ \d{4}\b', text)
    dates = [normalize_date(d) for d in dates]

    return {
        'name': name,
        'email': email_match.group(0) if email_match else None,
        'phone': phone_match.group(0) if phone_match else None,
        'skills': skills,
        'dates': dates,
        'raw_text': text
    }

